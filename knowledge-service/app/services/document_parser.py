import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.errors import (
    DocumentParsingError,
    EmptyDocumentError,
    UnsupportedDocumentError,
)


@dataclass(frozen=True)
class ParsedPage:
    page: int
    text: str


class DocumentParser:
    supported_extensions = {".pdf", ".docx"}

    def parse(self, filename: str, content: bytes) -> list[ParsedPage]:
        extension = Path(filename).suffix.lower()
        if extension not in self.supported_extensions:
            raise UnsupportedDocumentError("仅支持 PDF 和 DOCX 文件")
        if not content:
            raise EmptyDocumentError("上传文件不能为空")

        try:
            pages = (
                self._parse_pdf(content)
                if extension == ".pdf"
                else self._parse_docx(content)
            )
        except (UnsupportedDocumentError, EmptyDocumentError):
            raise
        except Exception as exception:
            raise DocumentParsingError(f"文件解析失败: {filename}") from exception

        cleaned_pages = [
            ParsedPage(page=item.page, text=self.clean_text(item.text))
            for item in pages
            if self.clean_text(item.text)
        ]
        if not cleaned_pages:
            raise EmptyDocumentError("文件中没有可索引的文本")
        return cleaned_pages

    def _parse_pdf(self, content: bytes) -> list[ParsedPage]:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exception:
                raise DocumentParsingError("暂不支持有密码保护的 PDF") from exception
        return [
            ParsedPage(page=index, text=page.extract_text() or "")
            for index, page in enumerate(reader.pages, start=1)
        ]

    def _parse_docx(self, content: bytes) -> list[ParsedPage]:
        document = DocxDocument(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
        # DOCX does not expose stable rendered page boundaries without a layout
        # engine, so all extracted content is explicitly assigned to page 1.
        return [ParsedPage(page=1, text="\n".join(blocks))]

    @staticmethod
    def clean_text(text: str) -> str:
        text = text.replace("\x00", " ").replace("\r\n", "\n")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r" *\n *", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
