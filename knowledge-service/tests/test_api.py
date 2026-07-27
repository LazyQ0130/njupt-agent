from io import BytesIO

import chromadb
from chromadb.config import Settings as NativeChromaSettings
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.core.config import (
    ChromaSettings,
    ChunkingSettings,
    EmbeddingSettings,
    Settings,
)
from app.main import create_app


def build_docx() -> bytes:
    output = BytesIO()
    document = DocxDocument()
    document.add_heading("南京邮电大学本科生转专业管理办法", level=1)
    document.add_paragraph(
        "学生申请转专业应关注教务处当年度通知，并满足目标学院公布的接收条件。"
        "学院可通过笔试、面试等方式组织考核。"
    )
    document.save(output)
    return output.getvalue()


def test_index_then_search_returns_source(tmp_path) -> None:
    settings = Settings(
        embedding=EmbeddingSettings(provider="hash", hash_dimensions=128),
        chroma=ChromaSettings(
            persist_directory=tmp_path / "chroma",
            collection_name="njupt_test_collection",
        ),
        chunking=ChunkingSettings(
            chunk_size_tokens=500,
            chunk_overlap_tokens=50,
        ),
        default_top_k=3,
    )

    with TestClient(create_app(settings)) as client:
        index_response = client.post(
            "/documents/index",
            files={
                "file": (
                    "本科生转专业管理办法.docx",
                    build_docx(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            data={
                "source": "教务处",
                "upload_time": "2026-07-25T12:00:00+08:00",
            },
        )
        search_response = client.post(
            "/rag/search",
            json={"question": "南邮转专业需要什么条件"},
        )

    assert index_response.status_code == 201
    assert index_response.json()["chunks_indexed"] >= 1
    assert index_response.json()["category"] == "ACADEMIC"
    assert search_response.status_code == 200
    assert search_response.json()
    assert search_response.json()[0]["filename"] == "本科生转专业管理办法.docx"
    assert search_response.json()[0]["page"] == 1
    assert search_response.json()[0]["source"] == "教务处"
    assert "转专业" in search_response.json()[0]["content"]
    assert 0 <= search_response.json()[0]["score"] <= 1


def test_rejects_legacy_doc_format(tmp_path) -> None:
    settings = Settings(
        embedding=EmbeddingSettings(provider="hash", hash_dimensions=64),
        chroma=ChromaSettings(
            persist_directory=tmp_path / "chroma",
            collection_name="njupt_doc_validation",
        ),
    )

    with TestClient(create_app(settings)) as client:
        response = client.post(
            "/documents/index",
            files={"file": ("legacy.doc", b"not-a-docx", "application/msword")},
            data={"source": "教务处"},
        )

    assert response.status_code == 422
    assert response.json()["detail"] == "仅支持 PDF 和 DOCX 文件"


def test_indexes_web_text_and_replaces_same_source_url(tmp_path) -> None:
    collection_name = "njupt_web_collection"
    persist_directory = tmp_path / "chroma"
    settings = Settings(
        embedding=EmbeddingSettings(provider="hash", hash_dimensions=128),
        chroma=ChromaSettings(
            persist_directory=persist_directory,
            collection_name=collection_name,
        ),
        chunking=ChunkingSettings(
            chunk_size_tokens=500,
            chunk_overlap_tokens=50,
        ),
    )
    base_payload = {
        "title": "本科生转专业通知",
        "source": "南京邮电大学本科生院",
        "source_url": (
            "https://jwc.njupt.edu.cn/2026/0701/c1594a300001/page.htm"
        ),
        "category": "ACADEMIC",
        "crawl_time": "2026-07-26T12:00:00+08:00",
        "last_updated": "2026-07-01T00:00:00+08:00",
    }

    with TestClient(create_app(settings)) as client:
        first = client.post(
            "/documents/index-text",
            json={
                **base_payload,
                "content": (
                    "旧版转专业通知要求学生提交纸质材料。"
                    "本段仅用于验证网页更新后旧向量会被替换。"
                    "申请人还需要在规定时间内完成学院审核和材料确认。"
                ),
                "content_hash": "a" * 64,
            },
        )
        second = client.post(
            "/documents/index-text",
            json={
                **base_payload,
                "content": (
                    "新版转专业通知要求学生关注当年度接收计划和目标学院考核，"
                    "具体申请时间以本科生院最新通知为准。"
                    "学生完成网上申请后，应继续关注审核结果和学校公示。"
                ),
                "content_hash": "b" * 64,
            },
        )
        search = client.post(
            "/rag/search",
            json={"question": "转专业接收计划和学院考核", "top_k": 1},
        )

    assert first.status_code == 201
    assert second.status_code == 201
    assert search.status_code == 200
    result = search.json()[0]
    assert result["filename"] == "本科生转专业通知"
    assert result["source_url"] == base_payload["source_url"]
    assert result["category"] == "ACADEMIC"
    assert "新版转专业通知" in result["content"]
    collection = chromadb.PersistentClient(
        path=str(persist_directory),
        settings=NativeChromaSettings(anonymized_telemetry=False),
    ).get_collection(collection_name)
    assert collection.count() == 1


def test_deletes_document_vectors_idempotently(tmp_path) -> None:
    collection_name = "njupt_delete_collection"
    persist_directory = tmp_path / "chroma"
    settings = Settings(
        embedding=EmbeddingSettings(provider="hash", hash_dimensions=128),
        chroma=ChromaSettings(
            persist_directory=persist_directory,
            collection_name=collection_name,
        ),
    )
    payload = {
        "title": "人工核验的缓考办理知识",
        "content": (
            "结论：学生因病无法参加考试时，应在考试前申请缓考。"
            "流程：准备证明材料，提交学院审核，并等待本科生院审批。"
            "未获批准而缺考的，按照学校考试管理规定处理。"
        ),
        "source": "南京邮电大学本科生院",
        "source_url": "https://jwc.njupt.edu.cn/rules/exam/page.htm",
        "category": "ACADEMIC",
        "crawl_time": "2026-07-27T12:00:00+08:00",
        "last_updated": "2026-07-27T12:00:00+08:00",
        "content_hash": "c" * 64,
        "source_type": "CURATED_OFFICIAL",
    }

    with TestClient(create_app(settings)) as client:
        indexed = client.post("/documents/index-text", json=payload)
        document_id = indexed.json()["document_id"]
        first_delete = client.delete(f"/documents/{document_id}")
        second_delete = client.delete(f"/documents/{document_id}")
        search = client.post(
            "/rag/search",
            json={"question": "因病如何申请缓考", "top_k": 3},
        )

    assert indexed.status_code == 201
    assert first_delete.status_code == 204
    assert second_delete.status_code == 204
    assert search.status_code == 200
    assert search.json() == []


def test_indexes_all_static_categories_into_vector_metadata(tmp_path) -> None:
    collection_name = "njupt_static_categories"
    persist_directory = tmp_path / "chroma"
    settings = Settings(
        embedding=EmbeddingSettings(provider="hash", hash_dimensions=128),
        chroma=ChromaSettings(
            persist_directory=persist_directory,
            collection_name=collection_name,
        ),
    )
    categories = {
        "SCHOOL_OVERVIEW": "学校章程规定办学宗旨、治理结构和师生权利义务。",
        "ORGANIZATION": "教学机构包括各学院，党政群部门依照职责服务师生。",
        "RESEARCH": "科研平台包括重点实验室、研究中心和学科建设基地。",
    }

    with TestClient(create_app(settings)) as client:
        for index, (category, sentence) in enumerate(categories.items(), start=1):
            response = client.post(
                "/documents/index-text",
                json={
                    "title": f"静态知识页面{index}",
                    "content": sentence * 4,
                    "source": "南京邮电大学官网",
                    "source_url": (
                        f"https://www.njupt.edu.cn/static/{index}/list.htm"
                    ),
                    "category": category,
                    "crawl_time": "2026-07-27T12:00:00+08:00",
                    "last_updated": "2026-07-27T12:00:00+08:00",
                    "content_hash": str(index) * 64,
                },
            )
            assert response.status_code == 201
            assert response.json()["category"] == category

    collection = chromadb.PersistentClient(
        path=str(persist_directory),
        settings=NativeChromaSettings(anonymized_telemetry=False),
    ).get_collection(collection_name)
    metadatas = collection.get(include=["metadatas"])["metadatas"]
    assert {metadata["category"] for metadata in metadatas} == set(categories)
