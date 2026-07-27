from io import BytesIO

import pytest
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas

from app.core.errors import UnsupportedDocumentError
from app.services.document_parser import DocumentParser


def build_docx(text: str) -> bytes:
    output = BytesIO()
    document = DocxDocument()
    document.add_heading("南京邮电大学", level=1)
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "来源"
    table.cell(0, 1).text = "教务处"
    document.save(output)
    return output.getvalue()


def build_pdf(text: str) -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 760, text)
    pdf.showPage()
    pdf.save()
    return output.getvalue()


def test_parses_and_cleans_docx() -> None:
    pages = DocumentParser().parse(
        "transfer.docx",
        build_docx("转专业申请需要关注当年度通知。"),
    )

    assert len(pages) == 1
    assert pages[0].page == 1
    assert "转专业申请" in pages[0].text
    assert "教务处" in pages[0].text


def test_parses_pdf_with_page_metadata() -> None:
    pages = DocumentParser().parse(
        "guide.pdf",
        build_pdf("NJUPT transfer policy"),
    )

    assert len(pages) == 1
    assert pages[0].page == 1
    assert "NJUPT transfer policy" in pages[0].text


def test_rejects_unsupported_extension() -> None:
    with pytest.raises(UnsupportedDocumentError):
        DocumentParser().parse("guide.doc", b"legacy word")
