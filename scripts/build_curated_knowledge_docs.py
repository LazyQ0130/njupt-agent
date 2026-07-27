from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path(__file__).resolve().parents[1] / "output" / "curated_knowledge"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 100, 112)
BLACK = RGBColor(20, 26, 34)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"


def set_run_font(
    run,
    *,
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_sections(doc: Document, running_label: str) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.paragraph_format.space_after = Pt(0)
        set_run_font(hp.add_run(running_label), size=8.5, color=MUTED)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(0)
        set_run_font(fp.add_run("南京邮电大学官方资料人工筛选版 · 2026-07-27"), size=8, color=MUTED)


def add_cover(doc: Document, kicker: str, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(105)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(kicker), size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run(title), size=27, color=RGBColor(32, 55, 72), bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    set_run_font(p.add_run(subtitle), size=13.5, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(66)
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.add_run("筛选原则"), size=10, color=MUTED, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.7)
    p.paragraph_format.right_indent = Inches(0.7)
    p.paragraph_format.line_spacing = 1.35
    set_run_font(
        p.add_run("只保留学生高频需要、可操作、可核验的内容；新闻宣传、党建活动和教师日常动态不纳入。"),
        size=11,
        color=BLACK,
    )
    doc.add_page_break()


def add_callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, [9360], indent=120)
    cell = table.cell(0, 0)
    shade_cell(cell, "F4F6F9")
    set_cell_margins(cell, 140, 140, 180, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(f"{label}："), size=10.5, color=DARK_BLUE, bold=True)
    set_run_font(p.add_run(text), size=10.5, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_source(doc: Document, label: str, url: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(f"{label}："), size=9, color=MUTED, bold=True)
    set_run_font(p.add_run(url), size=9, color=MUTED)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths: list[int], *, indent: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell, 90, 90, 120, 120)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[int],
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        shade_cell(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=9.5, color=DARK_BLUE, bold=True)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, text in enumerate(values):
            cell = cells[col_index]
            if row_index % 2 == 1:
                shade_cell(cell, "FAFBFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(text), size=9.2, color=BLACK)

    set_table_geometry(table, widths, indent=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_common_guide() -> Path:
    doc = Document()
    configure_styles(doc)
    configure_sections(doc, "学生常用政策与办事知识")
    add_cover(
        doc,
        "NJUPT STUDENT KNOWLEDGE",
        "南邮学生常用政策与办事知识",
        "基于 2025 版《学生手册》及 2025—2026 年官方页面人工筛选",
    )

    doc.add_heading("1. 使用说明", level=1)
    add_callout(
        doc,
        "适用范围",
        "面向本科生为主，同时给出研究生和其他学生群体的官方入口。涉及具体日期、名额、金额、材料清单时，以当学期本科生院、学生工作部（处）、研究生院或所在学院最新通知为准。",
    )
    add_bullets(
        doc,
        [
            "这份知识只保留“学生需要做决定或办理事项时会用到”的内容。",
            "不收录普通新闻、教师活动、党建动态、宣传报道和过期活动通知。",
            "遇到规则与学院通知不一致时，先看学校制度，再看当期通知和学院实施细则。",
        ],
    )

    doc.add_heading("2. 学校与校区", level=1)
    doc.add_paragraph(
        "南京邮电大学是国家“双一流”建设高校，以电子信息为鲜明特色。学校现有仙林、三牌楼、锁金村三个校区；学生手册同时说明学校在浦口区设有产教融合学院。"
    )
    add_bullets(
        doc,
        [
            "仙林校区：南京市栖霞区文苑路 9 号，邮编 210023。",
            "三牌楼校区：南京市新模范马路 66 号，邮编 210003。",
            "锁金村校区：南京市龙蟠路 177 号，邮编 210042。",
            "学校主页底部提供校历、规章制度、服务流程、图书馆、VPN 和校区地图等常用入口。",
        ],
    )
    add_source(doc, "学校主页", "https://www.njupt.edu.cn/")
    add_source(doc, "2025 版学生手册", "https://xsc.njupt.edu.cn/_upload/article/files/1e/2c/af7879b84b10b39a320d210e6fb4/bd5ebfd5-563f-4f71-bb4f-dc12aa83471b.pdf")

    doc.add_heading("3. 选课、考试与学籍", level=1)
    doc.add_heading("3.1 选课", level=2)
    add_bullets(
        doc,
        [
            "本科生应按本专业培养方案选课。当学期必修课通常由本科生院预置；体育项目、限制性选修课和任选课仍需在规定阶段上网选择。",
            "选课通常分为预选、正式选课和补改选三个阶段。未在规定阶段完成选课，可能影响参加考试和成绩记载。",
            "跨出本专业培养方案加修课程，需要在规定时间提出申请，并按学院和本科生院要求办理。",
            "每学期具体时间、课程容量和操作入口以本科生院通知及教务系统为准。",
        ],
    )
    add_source(doc, "学生选读课程实施细则（2024 修订）", "https://jwc.njupt.edu.cn/_t3997/2024/0430/c1755a261603/page.htm")

    doc.add_heading("3.2 缓考、补考和重修", level=2)
    add_bullets(
        doc,
        [
            "因疾病、代表学校参加重大竞赛或活动、直系亲属病危或亡故等情况无法参加期末考试，可按规定申请缓考。",
            "除考试现场突发情况外，缓考手续应在考试前完成，并准备相应证明，经学院和本科生院审批。",
            "未获批准而缺考通常按旷考处理；缓考不及格不再安排补考，需要申请重修。",
            "补考、重修的适用课程、时间和成绩记载方式以 2025 版学生手册及当学期通知为准。",
        ],
    )

    doc.add_heading("3.3 转专业与学籍异动", level=2)
    add_bullets(
        doc,
        [
            "转专业由学生本人提出申请，所在学院、接收学院和学校按当期工作方案审核与考核。",
            "2025—2026 学年第二学期转专业通知面向 2024 级和 2025 级普通本科学生；具体年级、接收名额、专业要求和考核方式每次都会变化。",
            "转入专业拟接收人数、考核内容和材料清单以接收学院当期细则为准，不应使用往年通知替代。",
            "休学、复学、转学、编级等学籍异动，应先联系所在学院教务办公室，再按本科生院学籍管理要求办理。",
        ],
    )
    add_source(doc, "2026 年转专业工作通知", "https://jwc.njupt.edu.cn/2026/0416/c1594a300221/page.htm")
    add_source(doc, "本科生学籍管理办法（2024 修订）", "https://jwc.njupt.edu.cn/_t278/2024/0430/c1535a261606/page.htm")

    doc.add_heading("4. 奖助、困难认定与勤工助学", level=1)
    doc.add_paragraph(
        "学校学生资助体系包括校内奖学金、国家奖学金、国家励志奖学金、国家助学金、家庭经济困难学生认定、生源地信用助学贷款、勤工助学、临时困难补助和特殊困难学生学杂费减免等。"
    )
    add_numbered(
        doc,
        [
            "先看学生工作部（处）“资助工作”或“学生手册（2025 版）”，确认项目、对象和当年通知。",
            "需要家庭经济困难认定的学生，按学院通知提交材料并参加认定；认定结果是多项资助申请的重要基础。",
            "奖助申请通常由学院组织，学生应留意辅导员、学院和智慧学工平台通知。",
            "如遇突发家庭困难，可向辅导员或学生资助管理中心咨询临时困难补助与绿色通道。",
        ],
    )
    add_callout(
        doc,
        "学生资助管理中心",
        "负责奖学金、助学贷款、勤工助学、困难认定、助学金和学费减免等工作。办公地址：仙林校区学生事务中心 310 室；电话：025-85866357；邮箱：zizhu@njupt.edu.cn。",
    )
    add_source(doc, "学生工作部（处）", "https://xsc.njupt.edu.cn/")
    add_source(doc, "家庭经济困难学生认定办法", "https://xsc.njupt.edu.cn/2024/0301/c17378a256066/page.htm")

    doc.add_heading("5. 图书馆与学习资源", level=1)
    add_bullets(
        doc,
        [
            "入馆和借阅使用本人校园一卡通。图书馆实行开架阅览，借还手续由服务总台或自助设备办理。",
            "本科生、研究生等校内读者每证通常限借 30 册；外文原版图书、密集书库等资源有单独上限。",
            "仙林和三牌楼图书馆设有自助借还设备；还书操作不需要刷卡。",
            "图书馆官网提供馆藏查询、数字资源、开放时间、馆藏分布、论文提交、查收查引、科技查新和新生入馆教育入口。",
            "节假日和寒暑假开放安排会变化，应查看图书馆最新通知。",
        ],
    )
    add_source(doc, "图书馆主页", "https://lib.njupt.edu.cn/")
    add_source(doc, "借阅规则", "https://lib.njupt.edu.cn/1385/list.htm")
    add_source(doc, "自助服务", "https://lib.njupt.edu.cn/zzfw/list.htm")

    doc.add_heading("6. 宿舍、安全、医疗与心理支持", level=1)
    add_bullets(
        doc,
        [
            "宿舍管理、请假考勤、校园治安、交通、电动自行车、宿舍防火和学生医保规则均收录在 2025 版学生手册中。",
            "住宿、门禁、维修、宿舍调整等具体问题，优先联系所在宿舍管理人员、辅导员或后勤管理处；不要使用过期的学院新闻作为办理依据。",
            "学生医保政策、待遇期限和门诊安排可能随地方政策调整，应查看学校当年通知并咨询校门诊部或相关部门。",
            "心理健康教育与咨询中心面向师生提供心理辅导、心理咨询和心理健康教育。需要帮助时可通过心理中心或辅导员获取预约与转介方式。",
        ],
    )
    add_source(doc, "学生工作部心理健康入口", "https://xl.njupt.edu.cn/")
    add_source(doc, "后勤管理处", "https://hqc.njupt.edu.cn/")

    doc.add_heading("7. 就业、升学与研究生培养", level=1)
    add_bullets(
        doc,
        [
            "就业岗位、双选会、生涯咨询和就业手续应以南京邮电大学就业信息网及学院毕业班通知为准。",
            "本科推免、考研、毕业与学位事项应分别查看本科生院、研究生招生网和学院通知，避免将招生宣传新闻当成政策。",
            "研究生培养方案按年级发布，研究生应选择本人入学年级对应的硕士或博士培养方案。",
            "研究生考试、学籍、科研实践创新、学位论文评审和学位申请均有独立规章或当期通知，办理前应核对研究生院最新版本。",
        ],
    )
    add_source(doc, "就业信息网", "https://njupt.91job.org.cn/")
    add_source(doc, "研究生院", "https://pg.njupt.edu.cn/")
    add_source(doc, "研究生培养方案", "https://pg.njupt.edu.cn/_t562/967/list.htm")
    add_source(doc, "研究生招生网", "https://yzb.njupt.edu.cn/")

    doc.add_heading("8. 高频问题应该去哪里查", level=1)
    add_table(
        doc,
        ["问题", "首选官方入口", "判断原则"],
        [
            ["本学期选课、考试、校历", "本科生院 jwc.njupt.edu.cn", "看当学期通知，不套用往年时间"],
            ["转专业、休复学、学籍", "本科生院 + 所在学院教务办", "学校规则 + 当期学院细则"],
            ["奖学金、助学金、困难认定", "学生工作部 xsc.njupt.edu.cn", "看项目对象与当年申报通知"],
            ["图书借阅、开放时间、数据库", "图书馆 lib.njupt.edu.cn", "节假日先看最新开放通知"],
            ["就业、双选会、招聘", "njupt.91job.org.cn", "以就业网岗位和学院通知为准"],
            ["研究生培养与学位", "研究生院 pg.njupt.edu.cn", "按本人年级和培养类型查询"],
            ["学院专业、培养与院内通知", "学院官网", "先确认学院与专业，再看对应栏目"],
        ],
        [1900, 2800, 4660],
    )

    doc.add_heading("9. 不应直接入库的内容", level=1)
    add_bullets(
        doc,
        [
            "普通新闻报道、领导调研、党建活动、教师获奖和与学生办事无关的会议动态。",
            "已结束的讲座、比赛、招聘会和节假日安排，除非保留其长期有效的办理规则。",
            "没有明确发布日期且内容明显依赖时效的通知。",
            "与标题关键词偶然匹配、但正文并非学生政策或服务的页面。",
            "正文乱码、编码异常、来源非南邮官方 HTTPS 域名或无法核验的内容。",
        ],
    )

    path = OUTPUT_DIR / "南邮学生常用政策与办事知识_人工筛选版.docx"
    doc.save(path)
    return path


UNDERGRAD_ROWS = [
    ["通信与信息工程学院", "通信工程；电子信息工程；广播电视工程", "scie.njupt.edu.cn"],
    ["电子与光学工程学院、柔性电子（未来技术）学院", "电子科学与技术；电磁场与无线技术；光电信息科学与工程；柔性电子学", "eoe.njupt.edu.cn"],
    ["集成电路科学与工程学院（产教融合学院）", "微电子科学与工程；集成电路设计与集成系统；集成电路科学与工程", "ic.njupt.edu.cn"],
    ["计算机学院、软件学院、网络空间安全学院", "计算机科学与技术；信息安全；软件工程", "cs.njupt.edu.cn"],
    ["自动化学院", "自动化；测控技术与仪器；电气工程及其自动化；智能电网信息工程", "coa.njupt.edu.cn"],
    ["人工智能学院", "人工智能；数据科学与大数据技术；机器人工程", "ai.njupt.edu.cn"],
    ["材料科学与工程学院", "高分子材料与工程；材料物理；新能源材料与器件", "iam.njupt.edu.cn"],
    ["化学与生命科学学院", "材料化学；分子科学与工程；生物医学工程", "chem.njupt.edu.cn"],
    ["物联网学院", "网络工程；物联网工程；地理信息科学；测绘工程；智能感知工程", "ciot.njupt.edu.cn"],
    ["理学院", "信息与计算科学（含省拔尖基地）；应用统计学；应用物理学；应用物理学+电子科学与技术双学士项目", "lxy.njupt.edu.cn"],
    ["现代邮政学院、智慧交通学院", "物流管理；邮政工程；物流管理+人工智能双学士项目", "simp.njupt.edu.cn"],
    ["数字媒体与设计艺术学院", "数字媒体艺术；广告学；数字媒体技术", "cm.njupt.edu.cn"],
    ["管理学院", "工商管理；人力资源管理；财务管理；大数据管理与应用；信息管理与信息系统+人工智能双学士项目", "bc.njupt.edu.cn"],
    ["经济学院", "经济统计学；金融工程；金融科技；数字经济", "jjxy.njupt.edu.cn"],
    ["社会与人口学院、社会工作学院", "行政管理；社会工作", "sps.njupt.edu.cn"],
    ["外国语学院", "英语；翻译", "fld.njupt.edu.cn"],
    ["教育科学与技术学院", "教育技术学", "edu.njupt.edu.cn"],
    ["贝尔英才学院", "通信工程；电子科学与技术；计算机科学与技术；人工智能", "bhs.njupt.edu.cn"],
    ["波特兰学院", "通信工程（中外合作办学）；电子科学与技术（中外合作办学）", "psu.njupt.edu.cn"],
]


OTHER_UNITS = [
    ["思想政治教育", "马克思主义学院", "marxism.njupt.edu.cn"],
    ["国际学生教育", "海外教育学院", "overseas.njupt.edu.cn"],
    ["中外合作办学", "欧洲塞浦路斯学院", "eci.njupt.edu.cn"],
    ["体育教学", "体育部", "tyb.njupt.edu.cn"],
    ["继续教育", "继续教育学院、自学考试办公室", "jjy.njupt.edu.cn"],
    ["应用型人才培养", "应用技术学院", "yyjsxy.njupt.edu.cn"],
    ["工程实践教学", "工程实验教学部", "deei.njupt.edu.cn"],
]


def build_college_guide() -> Path:
    doc = Document()
    configure_styles(doc)
    configure_sections(doc, "学院与专业导航")
    add_cover(
        doc,
        "NJUPT COLLEGE GUIDE",
        "南邮学院与专业导航",
        "以 2025 年本科招生专业口径为主，辅以学校教学机构目录",
    )

    doc.add_heading("1. 使用说明", level=1)
    add_callout(
        doc,
        "口径",
        "学院名称和官网来自南京邮电大学“教学机构”目录；本科招生专业来自学校 2025 年 6 月发布的《南京邮电大学 2025 年各学院本科招生专业》。2026 年及以后招生专业可能调整，报考或转专业时必须核对当年招生简章和学院细则。",
    )
    add_bullets(
        doc,
        [
            "这份导航回答“某个专业属于哪个学院”“学院官网在哪里”“哪些学院承担本科招生”等高频问题。",
            "学院官网应优先用于查询培养方案、院内教务通知、竞赛科研、实习就业和联系方式。",
            "学院新闻、党建活动和教师日常动态不作为专业或培养政策依据。",
        ],
    )
    add_source(doc, "学校教学机构目录", "https://www.njupt.edu.cn/jxjg/list.htm")
    add_source(doc, "2025 各学院本科招生专业", "https://zs.njupt.edu.cn/2025/0622/c2585a286175/pagem.htm")

    doc.add_heading("2. 2025 年本科招生学院与专业", level=1)
    doc.add_paragraph(
        "下表按学校 2025 年公开招生专业整理。专业名称后的双学士、中外合作办学、拔尖基地等属性是选择培养路径时的重要条件，不能只按专业简称判断。"
    )
    add_table(doc, ["学院", "2025 年本科招生专业", "学院官网"], UNDERGRAD_ROWS, [3000, 4500, 1860])

    doc.add_heading("3. 其他教学机构与学生查询入口", level=1)
    doc.add_paragraph(
        "以下单位同属学校教学机构，但不一定以普通本科招生学院的方式出现在 2025 专业表中。学生在查询公共课程、继续教育、国际教育或工程实践时，可能需要访问这些单位。"
    )
    add_table(doc, ["主要用途", "教学机构", "官网"], OTHER_UNITS, [2200, 3800, 3360])

    doc.add_heading("4. 学院资料的筛选规则", level=1)
    add_table(
        doc,
        ["保留", "暂不保留"],
        [
            ["学院简介、专业设置、培养方案、培养方向", "普通新闻、会议动态、领导调研"],
            ["本科/研究生教务通知、学籍与毕业要求", "只对教师适用的科研或人事通知"],
            ["实验室开放、竞赛、科研训练和创新项目", "没有报名入口或已结束且无长期价值的活动"],
            ["实习就业、升学与校企合作中的学生机会", "仅宣传企业或获奖结果、没有学生行动信息的新闻"],
            ["奖助、评优、转专业、推免和办事流程", "标题命中关键词但正文与学生无关的页面"],
            ["学院办公室、教务办、学工办联系方式", "个人隐私信息或非公开联系方式"],
        ],
        [4680, 4680],
    )

    doc.add_heading("5. 学生查询学院信息的推荐顺序", level=1)
    add_numbered(
        doc,
        [
            "先在本导航中确认专业归属学院和学院官网。",
            "进入学院官网后，优先查看“学院概况/专业介绍/人才培养/本科生教育/研究生教育”等长期栏目。",
            "办理具体事项时，再查看学院最新通知及学校对应职能部门通知。",
            "涉及名额、日期、材料和考核方式时，以当期通知为准；往年页面只能用于理解流程。",
            "若学院官网与学校文件口径冲突，先联系学院教务办或学工办核实。",
        ],
    )

    doc.add_heading("6. 不同问题对应的学院入口", level=1)
    add_bullets(
        doc,
        [
            "专业培养、选课替代、课程认定：学院教务办公室 + 本科生院。",
            "研究生培养、导师与学位：学院研究生办公室 + 研究生院。",
            "奖助评优、宿舍、就业与学生事务：学院学工办公室/辅导员 + 学生工作部。",
            "竞赛、科研训练、实验室：学院本科生教育或科研平台栏目；确认是否面向学生、是否仍在申报期。",
            "转专业：学校当期转专业通知 + 接收学院当期实施细则。",
            "招生咨询：本科招生网 + 招生学院官网；中外合作项目还需查看合作办学项目说明。",
        ],
    )

    path = OUTPUT_DIR / "南邮学院与专业导航_2025招生口径_人工筛选版.docx"
    doc.save(path)
    return path


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    common = build_common_guide()
    colleges = build_college_guide()
    print(common)
    print(colleges)


if __name__ == "__main__":
    main()
